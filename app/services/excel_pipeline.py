"""주간 경매 엑셀 업로드 파이프라인 (PRD §4, §5, §8).

Sheet `경매전체데이터` → AuctionRecord (raw lake) → MarketSummary (grid aggregate).
Sheet `시세표_테이블` (header row 5) → VehiclePriceTable (car-code reference).

Memory: openpyxl read_only + N-row bulk_insert_mappings (no full DataFrame / ORM lists).
"""
import logging
import re
from datetime import datetime, timezone

import openpyxl

from app.extensions import db
from app.models import AuctionRecord, MarketSummary, VehiclePriceTable
from app.services.accident import is_accident_free
from app.services.awd_utils import normalize_awd
from app.services.car_code import build_car_code, extract_grade_gdetail
from app.services.fuel_utils import normalize_fuel
from app.services.hierarchy import ensure_hierarchy
from app.services.analysis_logic import is_active
from app.services.mileage import calculate_mileage_bin

RAW_SHEET = "경매전체데이터"
logger = logging.getLogger(__name__)
PRICE_SHEET = "시세표_테이블"
BATCH_SIZE = 1000

# 실데이터 컬럼명 기준. PRD의 제조사/모델명 별칭 허용.
COLUMN_ALIASES = {
    "auction_date": ["경매일"],
    "maker": ["제작사", "제조사"],
    "kind": ["차종"],
    "model_name": ["모델명"],
    "car_name": ["차명"],
    "car_year": ["연식"],
    "car_km": ["주행거리"],
    "fuel": ["연료"],
    "start_price": ["시작가(만원)", "시작가"],
    "hope_price": ["희망가(만원)", "희망가"],
    "hammer_price": ["낙찰가"],
    "imported": ["내수수출구분", "내수/수출"],
    "accident_detail": ["사고'A' 상세", "사고A 상세"],
    "car_option": ["차량옵션"],
}
REQUIRED = ["auction_date", "maker", "car_name", "car_year", "car_km", "hammer_price", "imported"]

GROUP_KEYS = [
    "maker", "model_name", "mdetail_name", "grade_name", "gdetail_name",
    "car_year", "fuel", "awd", "is_accident_free", "imported", "km_bin",
]


class ExcelValidationError(ValueError):
    pass


def _headers_of(df_or_headers):
    if hasattr(df_or_headers, "columns"):
        return [str(c).strip() for c in df_or_headers.columns]
    return [str(c).strip() if c is not None else "" for c in df_or_headers]


def _resolve_columns(df_or_headers):
    """Map logical field → actual column name present in the sheet."""
    headers = _headers_of(df_or_headers)
    mapping = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in headers:
                mapping[field] = alias
                break
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        readable = {f: COLUMN_ALIASES[f][0] for f in missing}
        raise ExcelValidationError(f"필수 컬럼이 누락되었습니다: {list(readable.values())}")
    return mapping


def _find_header(headers, needle):
    for h in headers:
        if needle in str(h):
            return h
    return None


def _is_na(val):
    if val is None:
        return True
    try:
        if val != val:  # NaN
            return True
    except TypeError:
        pass
    return False


def _to_int(val):
    if _is_na(val):
        return None
    try:
        return int(float(re.sub(r"[^0-9.\-]", "", str(val)) or 0))
    except (TypeError, ValueError):
        return None


def _clean_str(val):
    """NaN / 'nan' / 공백 → None."""
    if _is_na(val):
        return None
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null", "-"):
        return None
    return s


def _clean_km(val):
    i = _to_int(val)
    return i if i is not None else 0


def _cell(row, index_by_name, name):
    if not name:
        return None
    i = index_by_name.get(name)
    if i is None or i >= len(row):
        return None
    return row[i]


def validate_excel_schema(df_or_headers):
    _resolve_columns(df_or_headers)
    return True


def _utcnow():
    return datetime.now(timezone.utc)


def _expunge_ingest_instances():
    """Drop identity-map rows from this ingest without detaching request User/History."""
    from app.models import (
        VehicleGrade, VehicleGradeDetail, VehicleMaker, VehicleModel, VehicleModelDetail,
    )
    kinds = (
        AuctionRecord, VehiclePriceTable, MarketSummary,
        VehicleMaker, VehicleModel, VehicleModelDetail, VehicleGrade, VehicleGradeDetail,
    )
    for obj in list(db.session.identity_map.values()):
        if isinstance(obj, kinds):
            db.session.expunge(obj)


def _flush_mappings(model, batch):
    if not batch:
        return
    db.session.bulk_insert_mappings(model, batch)
    db.session.commit()
    _expunge_ingest_instances()
    batch.clear()


def _row_to_auction_mapping(row, cols, index_by_name, week_no, xx_name, w_name):
    hammer = _to_int(_cell(row, index_by_name, cols["hammer_price"]))
    if hammer is None or hammer <= 0:
        return None
    maker = _clean_str(_cell(row, index_by_name, cols["maker"]))
    kind = _clean_str(_cell(row, index_by_name, cols.get("kind"))) if cols.get("kind") else None
    model_name = (
        _clean_str(_cell(row, index_by_name, cols.get("model_name")))
        if cols.get("model_name") else kind
    )
    car_name = _clean_str(_cell(row, index_by_name, cols["car_name"]))
    fuel = normalize_fuel(
        _cell(row, index_by_name, cols.get("fuel")) if cols.get("fuel") else None
    )
    imported = _clean_str(_cell(row, index_by_name, cols["imported"]))
    car_year = _to_int(_cell(row, index_by_name, cols["car_year"]))
    km = _clean_km(_cell(row, index_by_name, cols["car_km"]))
    option = (
        _clean_str(_cell(row, index_by_name, cols.get("car_option")))
        if cols.get("car_option") else None
    )
    awd = normalize_awd(car_name, option, None)
    acc_detail = (
        _clean_str(_cell(row, index_by_name, cols.get("accident_detail")))
        if cols.get("accident_detail") else None
    )
    xx = _cell(row, index_by_name, xx_name) if xx_name else None
    w = _cell(row, index_by_name, w_name) if w_name else None
    if _is_na(xx):
        xx = None
    if _is_na(w):
        w = None
    acc_free = is_accident_free(acc_detail, xx, w)

    model = kind or model_name
    mdetail = model_name or model
    grade, gdetail = extract_grade_gdetail(
        maker=maker, model=model, mdetail=mdetail, car_name=car_name,
    )
    hier = ensure_hierarchy(
        maker=maker, model=model, mdetail=mdetail, grade=grade, gdetail=gdetail,
    )
    code = build_car_code(
        maker=hier["maker"], model=hier["model"], mdetail=hier["mdetail"],
        grade=hier["grade"], gdetail=hier["gdetail"],
        car_year=car_year, fuel=fuel, awd=awd,
        car_name=car_name, car_option=option, accident_free=acc_free,
    )
    return {
        "week_no": week_no,
        "auction_date": _clean_str(_cell(row, index_by_name, cols["auction_date"])),
        "maker": hier["maker"],
        "model_name": hier["model"],
        "mdetail_name": hier["mdetail"],
        "grade_name": hier["grade"],
        "gdetail_name": hier["gdetail"],
        "car_name": car_name,
        "car_year": car_year,
        "car_km": km,
        "fuel": fuel,
        "awd": awd,
        "imported": imported,
        "start_price": (
            _to_int(_cell(row, index_by_name, cols.get("start_price")))
            if cols.get("start_price") else None
        ),
        "hope_price": (
            _to_int(_cell(row, index_by_name, cols.get("hope_price")))
            if cols.get("hope_price") else None
        ),
        "hammer_price": hammer,
        "accident_detail": acc_detail,
        "xx_exchange": _clean_str(xx) if xx is not None else None,
        "w_panel": _clean_str(w) if w is not None else None,
        "is_accident_free": acc_free,
        "maker_no": hier["maker_no"],
        "model_no": hier["model_no"],
        "mdetail_no": hier["mdetail_no"],
        "grade_no": hier["grade_no"],
        "gdetail_no": hier["gdetail_no"],
        "car_code": code,
        "km_bin": calculate_mileage_bin(km),
        "created_at": _utcnow(),
    }


def _ingest_raw_sheet(ws, week_no):
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise ExcelValidationError(f"시트 '{RAW_SHEET}' 가 비어 있습니다.")
    headers = _headers_of(header_row)
    cols = _resolve_columns(headers)
    index_by_name = {name: i for i, name in enumerate(headers)}
    xx_name = _find_header(headers, "XX 교환")
    w_name = _find_header(headers, "W 판금")

    batch = []
    inserted = 0
    for row in rows:
        mapping = _row_to_auction_mapping(
            row, cols, index_by_name, week_no, xx_name, w_name,
        )
        if mapping is None:
            continue
        batch.append(mapping)
        if len(batch) >= BATCH_SIZE:
            inserted += len(batch)
            _flush_mappings(AuctionRecord, batch)
    if batch:
        inserted += len(batch)
        _flush_mappings(AuctionRecord, batch)
    return inserted


def _ingest_price_sheet(ws):
    """시세표_테이블 header at row 5 → VehiclePriceTable via bulk mappings."""
    it = ws.iter_rows(values_only=True)
    headers = None
    name_col = fuel_col = imp_col = None
    year_cols = []
    index_by_name = {}
    batch = []
    inserted = 0
    for i, row in enumerate(it, start=1):
        if i < 5:
            continue
        if i == 5:
            headers = _headers_of(row)
            name_col = next((c for c in headers if "차명" in c), None)
            if name_col is None:
                return 0
            fuel_col = next((c for c in headers if "연료" in c), None)
            imp_col = next((c for c in headers if "내수" in c or "수출" in c), None)
            year_cols = [c for c in headers if re.fullmatch(r"20\d{2}", str(c))]
            index_by_name = {name: j for j, name in enumerate(headers)}
            continue
        name = _cell(row, index_by_name, name_col)
        if _is_na(name) or not str(name).strip():
            continue
        car_name = str(name).strip()
        fuel = _clean_str(_cell(row, index_by_name, fuel_col)) if fuel_col else None
        imported = _clean_str(_cell(row, index_by_name, imp_col)) if imp_col else None
        for yc in year_cols:
            price = _to_int(_cell(row, index_by_name, yc))
            if price is None or price == 0:
                continue
            batch.append({
                "car_name": car_name,
                "fuel": fuel,
                "imported": imported,
                "car_year": int(yc),
                "avg_price": price,
            })
            if len(batch) >= BATCH_SIZE:
                inserted += len(batch)
                _flush_mappings(VehiclePriceTable, batch)
    if batch:
        inserted += len(batch)
        _flush_mappings(VehiclePriceTable, batch)
    return inserted


def rebuild_market_summary(week_no):
    """Aggregate AuctionRecord into MarketSummary — 낙찰가 기준 · 전주대비(%). SQL GROUP BY."""
    if not is_active("aggregate.market_summary"):
        logger.info("rebuild_market_summary skipped: aggregate.market_summary inactive")
        return 0
    db.session.execute(db.delete(MarketSummary))
    db.session.commit()

    hammer_ok = (
        AuctionRecord.hammer_price.isnot(None),
        AuctionRecord.hammer_price > 0,
    )
    weeks_sorted = sorted(
        w for (w,) in db.session.execute(
            db.select(AuctionRecord.week_no).where(*hammer_ok).distinct()
        )
        if w is not None
    )
    if not weeks_sorted:
        return 0

    current_week = week_no if week_no in set(weeks_sorted) else weeks_sorted[-1]
    prev_week = None
    if current_week in weeks_sorted:
        idx = weeks_sorted.index(current_week)
        if idx > 0:
            prev_week = weeks_sorted[idx - 1]
    elif len(weeks_sorted) >= 2:
        current_week = weeks_sorted[-1]
        prev_week = weeks_sorted[-2]

    group_cols = [getattr(AuctionRecord, k) for k in GROUP_KEYS]

    def _agg_stmt(for_week):
        return (
            db.select(
                *group_cols,
                db.func.avg(AuctionRecord.hammer_price).label("hammer_avg"),
                db.func.avg(AuctionRecord.start_price).label("start_avg"),
                db.func.count(AuctionRecord.hammer_price).label("sample_count"),
                db.func.min(AuctionRecord.car_code).label("car_code"),
                db.func.min(AuctionRecord.car_name).label("car_name"),
            )
            .where(*hammer_ok, AuctionRecord.week_no == for_week)
            .group_by(*group_cols)
        )

    prev_lookup = {}
    if prev_week is not None:
        for row in db.session.execute(_agg_stmt(prev_week)).yield_per(BATCH_SIZE):
            m = row._mapping
            prev_lookup[tuple(m[k] for k in GROUP_KEYS)] = m["hammer_avg"]

    note = f"주간 집계 완료 ({current_week})"
    if prev_week:
        note += f" · 전주대비 기준 {prev_week}"

    batch = []
    count = 0
    for row in db.session.execute(_agg_stmt(current_week)).yield_per(BATCH_SIZE):
        m = row._mapping
        key = tuple(m[k] for k in GROUP_KEYS)
        prev = prev_lookup.get(key)
        cur = m["hammer_avg"]
        wow = None
        if prev and prev != 0:
            wow = round((cur - prev) / prev * 100, 2)
        batch.append({
            "car_code": m["car_code"],
            "maker": m["maker"],
            "model_name": m["model_name"],
            "mdetail_name": m["mdetail_name"],
            "grade_name": m["grade_name"],
            "gdetail_name": m["gdetail_name"],
            "car_name": m["car_name"],
            "car_year": _safe_int(m["car_year"]),
            "fuel": m["fuel"],
            "awd": m["awd"],
            "imported": m["imported"],
            "is_accident_free": bool(m["is_accident_free"]),
            "km_bin": m["km_bin"],
            "start_avg": _safe_float(m["start_avg"]),
            "hammer_avg": _safe_float(cur),
            "mom_pct": wow,
            "sample_count": int(m["sample_count"]),
            "week_no": current_week,
            "note": note,
        })
        if len(batch) >= BATCH_SIZE:
            count += len(batch)
            _flush_mappings(MarketSummary, batch)
    if batch:
        count += len(batch)
        _flush_mappings(MarketSummary, batch)
    return count


def _safe_int(v):
    try:
        if _is_na(v):
            return None
        return int(v)
    except (TypeError, ValueError):
        return None


def _safe_float(v):
    try:
        if _is_na(v):
            return None
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def process_weekly_upload(file_path, week_no, mode="append"):
    """Returns dict(rows_ok, records, summaries). Raises ExcelValidationError."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if RAW_SHEET not in wb.sheetnames:
            raise ExcelValidationError(f"시트 '{RAW_SHEET}' 가 없습니다.")

        if mode == "reset":
            db.session.execute(db.delete(AuctionRecord))
            db.session.execute(db.delete(VehiclePriceTable))
            db.session.commit()
        elif mode == "overwrite":
            db.session.execute(db.delete(AuctionRecord).where(AuctionRecord.week_no == week_no))
            db.session.commit()

        if PRICE_SHEET in wb.sheetnames:
            if mode in ("reset", "overwrite"):
                db.session.execute(db.delete(VehiclePriceTable))
                db.session.commit()
            _ingest_price_sheet(wb[PRICE_SHEET])

        # Price-table 24만 행 ORM 로드는 하지 않음. parse 경로가 references를 쓰지 않음.
        n = _ingest_raw_sheet(wb[RAW_SHEET], week_no)
        summaries = rebuild_market_summary(week_no)
        return {"rows_ok": n, "records": n, "summaries": summaries}
    finally:
        wb.close()


def clear_all_auction_data(clear_history=True):
    """업로드된 경매/시세 데이터를 전부 삭제."""
    from app.models import UploadHistory, VehicleGrade, VehicleGradeDetail, VehicleMaker, VehicleModel, VehicleModelDetail

    n_rec = db.session.execute(db.delete(AuctionRecord)).rowcount or 0
    n_sum = db.session.execute(db.delete(MarketSummary)).rowcount or 0
    n_price = db.session.execute(db.delete(VehiclePriceTable)).rowcount or 0
    db.session.execute(db.delete(VehicleGradeDetail))
    db.session.execute(db.delete(VehicleGrade))
    db.session.execute(db.delete(VehicleModelDetail))
    db.session.execute(db.delete(VehicleModel))
    db.session.execute(db.delete(VehicleMaker))
    n_hist = 0
    if clear_history:
        n_hist = db.session.execute(db.delete(UploadHistory)).rowcount or 0
    db.session.commit()
    return {
        "ok": True,
        "deleted_records": n_rec,
        "deleted_summaries": n_sum,
        "deleted_price_rows": n_price,
        "deleted_history": n_hist,
    }


def delete_upload_by_history(history_id):
    """특정 업로드 이력의 주차 데이터를 삭제하고 시세를 재집계."""
    from app.models import UploadHistory

    hist = db.session.get(UploadHistory, history_id)
    if hist is None:
        return {"ok": False, "error": "업로드 이력을 찾을 수 없습니다."}
    week_no = hist.week_no
    n = (db.session.execute(db.delete(AuctionRecord).where(AuctionRecord.week_no == week_no)).rowcount or 0) if week_no else 0
    db.session.delete(hist)
    db.session.commit()
    remaining = db.session.execute(db.select(AuctionRecord).order_by(AuctionRecord.id.desc())).scalars().first()
    summaries = rebuild_market_summary(remaining.week_no if remaining else week_no)
    return {
        "ok": True,
        "deleted_records": n,
        "week_no": week_no,
        "summaries": summaries,
    }
