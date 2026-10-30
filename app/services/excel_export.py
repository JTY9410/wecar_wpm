"""도매시세·전체시세 Excel 내보내기."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services import market_query

GRID_HEADERS = [
    "카코드", "제조사", "모델", "세부모델", "등급", "세부등급", "차명",
    "년식", "유종", "AWD", "사고등급", "주행구간",
    "시작가평균", "낙찰가평균", "전주대비(%)", "표본수", "주차",
]

FULL_GRID_HEADERS = [
    "제조사", "모델", "세부모델", "등급", "세부등급", "년식", "유종",
    "주행구간", "사고여부", "낙찰가평균", "표본수",
]


def _style_headers(ws, headers):
    header_fill = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws):
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14


def _to_workbook(title, headers, row_values) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _style_headers(ws, headers)
    for i, values in enumerate(row_values, 2):
        for col, val in enumerate(values, 1):
            ws.cell(row=i, column=col, value=val)
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_grid_excel(args) -> io.BytesIO:
    rows = market_query.grid(
        maker=args.get("maker"),
        model_name=args.get("model_name"),
        mdetail_name=args.get("mdetail_name"),
        car_name=args.get("car_name") or args.get("q"),
        grade_name=args.get("grade_name"),
        gdetail_name=args.get("gdetail_name"),
        car_year=args.get("car_year"),
        fuel=args.get("fuel"),
        awd=args.get("awd"),
        q=args.get("q"),
        lang="ko",
    )
    values = [
        [
            r.get("car_code"),
            r.get("maker"),
            r.get("model_name"),
            r.get("mdetail_name"),
            r.get("grade_name"),
            r.get("gdetail_name"),
            r.get("car_name"),
            r.get("car_year"),
            r.get("fuel"),
            r.get("awd"),
            r.get("accident_grade") or r.get("is_accident_free"),
            r.get("km_bin"),
            r.get("start_avg"),
            r.get("hammer_avg"),
            r.get("wow_pct"),
            r.get("sample_count"),
            r.get("week_no") or "",
        ]
        for r in rows
    ]
    return _to_workbook("도매시세", GRID_HEADERS, values)


def export_full_grid_excel(args) -> io.BytesIO:
    data = market_query.full_grid(
        week_from=args.get("week_from"),
        week_to=args.get("week_to"),
        maker=args.get("maker"),
        model_name=args.get("model_name"),
        mdetail_name=args.get("mdetail_name") or args.get("car_name"),
        grade_name=args.get("grade_name"),
        gdetail_name=args.get("gdetail_name"),
        car_year=args.get("car_year"),
        fuel=args.get("fuel"),
        q=args.get("q"),
        lang="ko",
        limit=10000,
    )
    values = [
        [
            r.get("maker"),
            r.get("model_name"),
            r.get("mdetail_name"),
            r.get("grade_name"),
            r.get("gdetail_name"),
            r.get("car_year"),
            r.get("fuel"),
            r.get("km_bin"),
            r.get("is_accident_free"),
            r.get("hammer_avg"),
            r.get("sample_count"),
        ]
        for r in (data.get("rows") or [])
    ]
    return _to_workbook("전체시세", FULL_GRID_HEADERS, values)
